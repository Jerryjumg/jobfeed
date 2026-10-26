#!/usr/bin/env python3
"""jobfeed: pull open roles from public job-board APIs, filter them by my rules,
score them with jdfit, and write one spreadsheet to work from each day.

    ../.tools/venv/bin/python jobfeed.py            # fetch, filter, write Jobs_<date>.xlsx
    ../.tools/venv/bin/python jobfeed.py discover names.txt   # find which boards exist

Sources: Greenhouse, Lever and Ashby public posting APIs (companies.json), plus the
community new-grad and internship lists on GitHub (config.json).
"""
from __future__ import annotations

import html
from collections import Counter
import json
import re
import sys
import time
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "jdfit"))
import jdfit  # noqa: E402

CONFIG = json.loads((HERE / "config.json").read_text())
SEEN_PATH = HERE / "seen.json"
STATE_PATH = HERE / "state.json"
TODAY = datetime.now().strftime("%Y-%m-%d")
OUT_PATH = HERE / f"Jobs_{TODAY}.xlsx"


def last_pull_date() -> str:
    """The day of the previous pull. Only postings from that day onward, never shown before, make the sheet.

    That day is included on purpose: a job posted in the evening after I pulled would otherwise never appear.
    """
    state = json.loads(STATE_PATH.read_text()) if STATE_PATH.exists() else {}
    if state.get("last_run") == TODAY:                      # second pull on the same day: keep the same window
        previous = state.get("previous_run")
    else:
        previous = state.get("last_run")
        if not previous and SEEN_PATH.exists():
            previous = datetime.fromtimestamp(SEEN_PATH.stat().st_mtime).strftime("%Y-%m-%d")
        STATE_PATH.write_text(json.dumps({"last_run": TODAY, "previous_run": previous}))
    floor = (datetime.now() - timedelta(days=CONFIG["max_days_back"])).strftime("%Y-%m-%d")
    return max(previous or floor, floor)


OLDEST = last_pull_date()
OLDEST_AGE = (datetime.now() - datetime.strptime(OLDEST, "%Y-%m-%d")).days

STATES = ("AL|AK|AZ|AR|CA|CO|CT|DE|FL|GA|HI|ID|IL|IN|IA|KS|KY|LA|ME|MD|MA|MI|MN|MS|MO|MT|NE|NV|NH|NJ|"
          "NM|NY|NC|ND|OH|OK|OR|PA|RI|SC|SD|TN|TX|UT|VT|VA|WA|WV|WI|WY|DC")
US_RE = re.compile(r"United States|\bUSA?\b|U\.S\.|" + "|".join(CONFIG["us_cities"]), re.I)
STATE_CODE_RE = re.compile(rf"(?:^|[,/ -])\s*(?:{STATES})\b")          # case-sensitive on purpose
STATE_RE = re.compile(rf"[A-Z][a-z]+(?: [A-Z][a-z]+)?,\s*(?:{STATES})\b")
NON_US_RE = re.compile("|".join(CONFIG["non_us_markers"]), re.I)
CANADA_RE = re.compile(r"[A-Za-z .]+,\s*(?:ON|BC|QC|AB|MB|NS)\b(?:,\s*CA(?:N|NADA)?\b)?|Canada", re.I)
HOME_RE = re.compile("|".join(CONFIG["home_markers"]), re.I)
HUB_RE = re.compile("|".join(CONFIG["odds"]["hub_locations"]), re.I)
COMPANY_OUT = re.compile("|".join(CONFIG["company_exclude"]), re.I)
TITLE_IN = re.compile("|".join(CONFIG["title_include"]), re.I)
TITLE_OUT = re.compile("|".join(CONFIG["title_exclude"]), re.I)
LANES = [(lane["name"], re.compile(lane["title"], re.I), lane["people"]) for lane in CONFIG["lanes"]]


def lane_of(title: str) -> int:
    """Index of the first lane whose pattern matches the title; the last lane is the catch-all."""
    for i, (_, pattern, _) in enumerate(LANES[:-1]):
        if pattern.search(title):
            return i
    return len(LANES) - 1


def title_wanted(title: str) -> bool:
    return bool(TITLE_IN.search(title)) or lane_of(title) < len(LANES) - 1


HARD = [(re.compile(p, re.I), why) for p, why in CONFIG["hard_exclusions"]]
VISA_SENTENCE = re.compile(r"(?i:sponsorship|(visa|immigration|employment|work)[ -]sponsor|will sponsor|\bvisas?\b|"
                           r"work authori[sz]|authori[sz]ed to work|citizen|green card|permanent resident|h-?1b|e-verify)"
                           r"|\bOPT\b|\bCPT\b(?! codes?)|\bF-?1\b")
EEO_RE = re.compile(r"without regard|regardless of|protected (by|under|class|veteran|characteristic)|equal opportunity|"
                    r"discriminat|national origin", re.I)
IGNORE_SENTENCES = re.compile("|".join(CONFIG["boilerplate_to_ignore"]), re.I)


def get(url: str, timeout: int = 60):
    req = urllib.request.Request(url, headers={"User-Agent": "jobfeed/1.0 (personal job search)"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


def to_text(raw: str) -> str:
    t = html.unescape(raw or "")
    t = re.sub(r"<li[^>]*>", "\n* ", t)
    t = re.sub(r"</(p|div|h\d|ul|ol|li)>|<br ?/?>", "\n", t)
    t = re.sub(r"<[^>]+>", " ", t)
    t = html.unescape(t)
    t = re.sub(r"[ \t\xa0]+", " ", t)
    return re.sub(r"\n\s*\n+", "\n", t).strip()


def day(value) -> str:
    if not value:
        return ""
    if isinstance(value, (int, float)):
        if value > 1e11:
            value /= 1000
        return datetime.fromtimestamp(value, timezone.utc).strftime("%Y-%m-%d")
    return str(value)[:10]


# --------------------------------------------------------------------------- sources
def greenhouse(slug: str) -> list[dict]:
    data = get(f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs?content=true")
    return [dict(company=slug, title=j["title"].strip(), location=j.get("location", {}).get("name", ""),
                 url=j["absolute_url"], posted=day(j.get("first_published") or j.get("updated_at")),
                 text=to_text(j.get("content", "")), source="greenhouse") for j in data.get("jobs", [])]


def lever(slug: str) -> list[dict]:
    out = []
    for j in get(f"https://api.lever.co/v0/postings/{slug}?mode=json"):
        body = (j.get("descriptionPlain") or "") + "\n" + "\n".join(
            f"{sec.get('text', '')}\n{to_text(sec.get('content', ''))}" for sec in j.get("lists", []))
        out.append(dict(company=slug, title=j["text"].strip(), location=j.get("categories", {}).get("location", ""),
                        url=j["hostedUrl"], posted=day(j.get("createdAt")), text=body, source="lever"))
    return out


def ashby(slug: str) -> list[dict]:
    data = get(f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=true")
    out = []
    for j in data.get("jobs", []):
        if not j.get("isListed", True):
            continue
        loc = "; ".join([j.get("location") or ""] + [s.get("location", "") for s in j.get("secondaryLocations", [])])
        out.append(dict(company=slug, title=j["title"].strip(), location=loc, url=j["jobUrl"],
                        posted=day(j.get("publishedAt")), text=to_text(j.get("descriptionHtml", "")), source="ashby",
                        pay=(j.get("compensation") or {}).get("compensationTierSummary") or ""))
    return out


def github_list(entry: dict) -> list[dict]:
    cutoff = time.time() - entry["max_age_days"] * 86400
    out = []
    for j in get(entry["url"], timeout=120):
        if not (j.get("active") and j.get("is_visible", True)) or j.get("date_posted", 0) < cutoff:
            continue
        out.append(dict(company=j["company_name"], title=j["title"].strip(), location="; ".join(j.get("locations", [])),
                        url=j["url"], posted=day(j["date_posted"]), text="", source=entry["name"],
                        sponsorship=j.get("sponsorship", ""), category=j.get("category", ""),
                        terms=", ".join(j.get("terms", []))))
    return out


def hn_hiring(_: str) -> list[dict]:
    """This month's Hacker News "Who is hiring?" thread, keeping posts that mention Boston/Massachusetts or remote."""
    hits = get("https://hn.algolia.com/api/v1/search_by_date?query=%22Who%20is%20hiring%22"
               "&tags=story,author_whoishiring&hitsPerPage=5")["hits"]
    thread = next(h for h in hits if h["title"].startswith("Ask HN: Who is hiring?"))
    posted = day(thread["created_at"])
    out = []
    for c in get(f"https://hn.algolia.com/api/v1/items/{thread['objectID']}", timeout=120).get("children", []):
        text = to_text(c.get("text") or "")
        if not text:
            continue
        head = text.split("\n", 1)[0]
        if not re.search(r"\bboston\b|cambridge, ma|somerville|massachusetts|\bremote\b", head, re.I):
            continue
        parts = [p.strip() for p in head.split("|")]
        location = next((p for p in parts if re.search(r"boston|cambridge|remote|onsite|hybrid|, ma\b", p, re.I)), "")
        out.append(dict(company=parts[0][:60], title=head[:160], location=location or head[:80],
                        url=f"https://news.ycombinator.com/item?id={c['id']}", posted=posted, text=text, source="hn"))
    return out


def post(url: str, payload: dict, timeout: int = 40):
    req = urllib.request.Request(url, data=json.dumps(payload).encode(),
                                 headers={"Content-Type": "application/json", "User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


def posted_days(label: str) -> int:
    label = (label or "").lower()
    if "today" in label:
        return 0
    if "yesterday" in label:
        return 1
    m = re.search(r"(\d+)\+? day", label)
    return int(m.group(1)) + (1 if "+" in label else 0) if m else 99


def workday(entry: dict) -> list[dict]:
    """Workday has no full dump, so run each keyword search and keep recent, relevant titles."""
    # Most tenants live on <host>.myworkdayjobs.com; a few (Magna) use the myworkdaysite.com layout instead.
    base = entry.get("base") or f"https://{entry['host']}.myworkdayjobs.com"
    api = f"{base}/wday/cxs/{entry['tenant']}/{entry['site']}"
    found: dict[str, dict] = {}
    deadline = time.time() + 150            # a slow or rate-limiting site must not stall the whole run
    for term in CONFIG["workday_searches"]:
        if time.time() > deadline:
            break
        for offset in (0, 20):
            try:
                page = post(f"{api}/jobs", {"appliedFacets": {}, "limit": 20, "offset": offset, "searchText": term},
                            timeout=15)
            except Exception:
                break
            postings = page.get("jobPostings", [])
            for p in postings:
                if p.get("externalPath") and p.get("title"):
                    found.setdefault(p["externalPath"], p)
            if len(postings) < 20:
                break
    out = []
    for path, p in found.items():
        age = posted_days(p.get("postedOn", ""))
        title = p["title"].strip()
        if age > OLDEST_AGE or TITLE_OUT.search(title) or not title_wanted(title):
            continue
        if time.time() > deadline + 120:
            break
        try:
            info = get(f"{api}{path}", timeout=15)["jobPostingInfo"]
        except Exception:
            continue
        country = (info.get("country") or {}).get("descriptor", "")
        loc = "; ".join([info.get("location", "")] + info.get("additionalLocations", []))
        if "United States" in country and not US_RE.search(loc):
            loc += ", United States"
        out.append(dict(company=entry["name"], title=title, location=loc, country=country,
                        url=info.get("externalUrl") or f"{base}{entry.get('public_path', '/' + entry['site'])}{path}",
                        posted=(datetime.now() - timedelta(days=age)).strftime("%Y-%m-%d"),
                        text=to_text(info.get("jobDescription", "")), source="workday"))
    return out


FETCHERS = {"greenhouse": greenhouse, "lever": lever, "ashby": ashby, "workday": workday, "hn": hn_hiring}


# --------------------------------------------------------------------------- filter
# Two-letter country codes that are not also US state codes, and patterns like "Hamburg HH DE" where a
# foreign region code sits in front of a code that would otherwise read as a US state (DE, IN, CA, IL).
COUNTRY_CODE_RE = re.compile(r"(?:^|[\s,/-])(?:IE|GB|UK|FR|NL|ES|PL|CH|SE|BE|AT|DK|CZ|PT|RO|HU|SG|JP|CN|MX|BR|AU|NZ|KR|TW|HK|"
                             r"IT|NO|FI|GR|TR|ZA|AE|PH|TH|MY|VN|CR|CL|LU|BG|RS|HR|SK|UA|LT|EE|LV)(?:$|[\s,/-])"
                             r"|\b[A-Z]{2}[ -](?:DE|IN|CA|IL)\b(?![ -]?US)")
VAGUE_LOCATION_RE = re.compile(r"^\W*$|hybrid|remote|in.?office|on.?site|multiple|\d+ locations|distributed|flexible|"
                               r"anywhere|various|headquarters|\bHQ\b|\bN/?A\b|global|worldwide", re.I)


STATE_NAMES = ("Alabama|Alaska|Arizona|Arkansas|California|Colorado|Connecticut|Delaware|Florida|Georgia|Hawaii|Idaho|"
               "Illinois|Indiana|Iowa|Kansas|Kentucky|Louisiana|Maine|Maryland|Massachusetts|Michigan|Minnesota|"
               "Mississippi|Missouri|Montana|Nebraska|Nevada|New Hampshire|New Jersey|New Mexico|New York|"
               "North Carolina|North Dakota|Ohio|Oklahoma|Oregon|Pennsylvania|Rhode Island|South Carolina|"
               "South Dakota|Tennessee|Texas|Utah|Vermont|Virginia|Washington|West Virginia|Wisconsin|Wyoming")
STATE_NAME_RE = re.compile(rf"\b(?:{STATE_NAMES})\b")
CITY_ABBR_RE = re.compile(r"\b(?:SF|SFO|NYC|NY|SEA|BOS|ATX|ATL|CHI|DFW|PHX|DEN|PDX|SLC|RTP|SoCal|NorCal|Bay Area|AMER)\b")
US_TEXT_RE = re.compile(r"United States|\bU\.S\.|\bUSA?\b|401\(k\)|\$\d{2,3},\d{3}|\bE-Verify\b")
LOCATION_LINE_RE = re.compile(r"(?:available |job |work |office )?locations?\s*(?:\(s\))?\s*:\s*([^\n]{2,160})", re.I)


def place_verdict(loc: str):
    """True for a US place, False for a foreign one, None when the text is too vague to tell."""
    original = loc
    loc = CANADA_RE.sub(" ", loc)
    named = re.sub(r"remote", " ", loc, flags=re.I)           # "Remote" alone says nothing about the country
    if (US_RE.search(loc) or HOME_RE.search(named) or HUB_RE.search(loc) or STATE_NAME_RE.search(loc)
            or CITY_ABBR_RE.search(loc)):
        return True
    if original.strip() and not loc.strip(" ;,/-"):         # nothing left once Canadian places are removed
        return False
    if re.search(r"canada", original, re.I) or NON_US_RE.search(loc) or COUNTRY_CODE_RE.search(loc):
        return False
    if STATE_CODE_RE.search(loc):
        return True
    return None if VAGUE_LOCATION_RE.search(loc) else False   # an unrecognized specific place counts as foreign


def is_us(job: dict) -> bool:
    if job.get("country"):                                    # Workday states the country outright; trust it
        return "United States" in job["country"]
    verdict = place_verdict(job["location"])
    if verdict is not None:
        return verdict
    # Vague location ("Hybrid", "Remote", "In-Office", "N/A"): read where the posting itself says the job is.
    text = job["text"]
    for line in LOCATION_LINE_RE.findall(text[:6000]):
        inner = place_verdict(line)
        if inner is not None:
            return inner
    head = text[:3000]
    if STATE_RE.search(head) or US_TEXT_RE.search(text):
        return True
    return not NON_US_RE.search(head)                         # no evidence either way: keep it


def judge(job: dict, profile: dict, patterns) -> tuple[bool, str]:
    """Return (keep, reason). Also fills verdict, score, years and gaps on the job."""
    title = job["title"]
    if COMPANY_OUT.search(job["company"]):
        return False, "company: defense contractor"
    if TITLE_OUT.search(title):
        return False, "title: senior or outside my lanes"
    if not title_wanted(title):
        return False, "title: not a technical role"
    if not is_us(job):
        return False, "location: outside the US"
    if job["source"] not in FETCHERS:                       # GitHub list entry, no JD text
        if job["posted"] and job["posted"] < OLDEST:
            return False, "old: posted before my last pull"
        if job.get("category") and job["category"] in CONFIG["list_categories_exclude"]:
            return False, f"list category: {job['category']}"
        terms = [t for t in (job.get("terms") or "").split(", ") if t]
        if terms and all("summer" in t.lower() or "fall" in t.lower() for t in terms):
            return False, "internship term runs past my OPT end date"
        if job.get("sponsorship") in CONFIG["list_sponsorship_exclude"]:
            return False, f"list says: {job['sponsorship']}"
        job.update(verdict="Open the posting", score=None, years=None, gaps="", _report=None, _text="")
        return True, ""
    if job["posted"] and job["posted"] < OLDEST:
        return False, "old: posted before my last pull"
    text = IGNORE_SENTENCES.sub(" ", job["text"])
    for pattern, why in HARD:
        if pattern.search(text):
            return False, why
    # Whatever the posting says about visas or work authorization, shown in the sheet so I can read it myself.
    visa = [x.strip() for x in re.split(r"[.\n]", text)
            if len(x) < 600 and VISA_SENTENCE.search(x) and not EEO_RE.search(x)]
    job["visa_note"] = " | ".join(dict.fromkeys(visa))[:300]
    report = jdfit.assess(text, profile, patterns, source=job["url"])
    gaps = [h.skill for h in report.skills if h.status == "gap" and h.bucket == "required"]
    job.update(verdict=report.verdict, score=report.score, years=report.years_required, gaps=", ".join(gaps[:6]),
               _report=report, _text=text)
    if report.verdict == "Too senior":
        return False, f"asks for {report.years_required}+ years"
    if report.verdict.startswith("Skip"):
        return False, "red flag: " + ", ".join(report.red_flags)
    return True, ""


# --------------------------------------------------------------------------- odds
ODDS = CONFIG["odds"]
BRAND_RE = re.compile("|".join(ODDS["big_brands"]), re.I)
JUNIOR_RE = re.compile(r"(?i:junior|\bjr\b|associate|entry|early.?career|new.?grad|graduate|university|apprentice)|\b(I|1)\b")
LEVEL2_RE = re.compile(r"\b(II|2)\b")
INTERN_RE = re.compile(r"intern\b|internship|co-?op", re.I)
EDGE_RES = [(re.compile(p, re.I), label, mult) for p, label, mult in ODDS["edges"]]
SPONSOR_YES = re.compile(r"(we|will|do|does|can|able to|happy to) sponsor|sponsorship (is )?(available|offered|provided)(?!\s*\?)|"
                         r"offers sponsorship", re.I)
SPONSOR_SOFT = re.compile(r"discretion|may not be able|not guaranteed|case[- ]by[- ]case|limited", re.I)
WEST_ONLY_RE = re.compile(ODDS.get("west_only_locations", r"$^"), re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[a-z]{2,}(?:\.[a-z]{2,})?", re.I)
NOREPLY_RE = re.compile(r"no-?reply|accommodat|privacy|eeo|benefits|ada@|disab", re.I)
ATS_SOURCES = ("greenhouse", "lever", "ashby")


def company_key(name: str) -> str:
    """Loose company name for matching a job board against LinkedIn's Company column."""
    name = re.sub(r"\(.*?\)|,? (inc|llc|ltd|corp|corporation|co|company|group|holdings|technologies|technology|labs?)\.?$",
                  "", name.lower().strip())
    return re.sub(r"[^a-z0-9]", "", name)


def load_connections() -> dict[str, list[str]]:
    """My LinkedIn connections by company, from the Connections.csv in LinkedIn's data export, if I saved one."""
    import csv
    for path in (HERE.parent / "Connections.csv", HERE.parent / "LinkedIn" / "Connections.csv"):
        if path.exists():
            lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
            start = next((i for i, l in enumerate(lines) if l.startswith("First Name")), None)
            if start is None:
                return {}
            people: dict[str, list[str]] = {}
            for row in csv.DictReader(lines[start:]):
                key = company_key(row.get("Company") or "")
                if key:
                    people.setdefault(key, []).append(
                        f"{row.get('First Name', '')} {row.get('Last Name', '')} ({row.get('Position', '')})".strip())
            return people
    return {}


CONNECTIONS = load_connections()


def contacts_for(job: dict, text: str) -> list[str]:
    """People I could write to about this job: an address in the posting, then anyone I'm connected to there."""
    found = []
    emails = [e for e in dict.fromkeys(EMAIL_RE.findall(text or "")) if not NOREPLY_RE.search(e)]
    if emails:
        found.append("email in posting: " + ", ".join(emails[:2]))
    known = CONNECTIONS.get(company_key(job["company"]))
    if known:
        found.append(f"you know {len(known)} there: " + "; ".join(known[:2]))
    return found


def odds(job: dict, report=None, text: str = "") -> tuple[int, str]:
    """A relative score for how likely this application is to turn into an offer.

    It multiplies the stages an application has to survive: being eligible, clearing the
    experience screen, matching the skills, and then beating the other applicants. It is not
    a measured probability; the weights live in config.json under "odds" so they can be tuned
    against real outcomes later.
    """
    f: dict[str, float] = {}
    title, loc = job["title"], job["location"]
    w = ODDS["weights"]

    f[LANES[job["lane"]][0][2:]] = ODDS["lane_prior"][job["lane"]]

    # 1. Experience screen
    years = job.get("years")
    if years is not None and years >= 2:
        f[f"asks {years}+ yrs"] = w["two_years"]
    elif years is not None:
        f[f"asks {years} yr"] = 1.0
    elif JUNIOR_RE.search(title) or job["source"] not in FETCHERS:
        f["junior title"] = 1.0
    else:
        f["no level stated"] = w["no_level_stated"]
    if LEVEL2_RE.search(title):
        f["level II title"] = w["level_two_title"]
    board = job.get("board_size") or 0
    small = job["source"] in ATS_SOURCES and 0 < board <= 20
    if INTERN_RE.search(title):
        # Small companies often take graduates for co-ops; large ones enforce enrollment.
        f["co-op at a small company" if small else "internship, may need enrollment"] = (
            w["internship_small"] if small else w["internship"])

    # 2. Skill match, shrunk toward the middle when the posting names few skills I can recognize
    if report is not None:
        hits = report.skills
        total = sum(3 if h.bucket == "required" else 1 for h in hits)
        matched = sum(3 if h.bucket == "required" else 1 for h in hits if h.status == "have")
        shrunk = (matched + 1.5) / (total + 3)
        gaps = sum(1 for h in hits if h.status == "gap" and h.bucket == "required")
        f[f"skills {round(shrunk * 100)}%"] = (0.4 + 0.6 * shrunk) * (w["required_gap"] ** min(gaps, 3))
    else:
        f["JD not read"] = w["no_jd"]

    # 3. Sponsorship signal
    visa = " ".join(x for x in (job.get("visa_note"), job.get("sponsorship")) if x)
    if SPONSOR_YES.search(visa):
        f["says it sponsors"] = w["sponsors"]
    elif SPONSOR_SOFT.search(visa):
        f["sponsorship uncertain"] = w["sponsorship_soft"]

    # 4. Competition: how many other people are likely to apply, and whether I have an edge on them
    if BRAND_RE.search(job["company"]):
        f["famous company"] = w["big_brand"]
    elif small:
        f[f"small company ({board} openings)"] = w["small_company"]
    elif job["source"] in ATS_SOURCES and board <= 60:
        f[f"mid-size company ({board} openings)"] = w["mid_company"]
    elif job["source"] in ATS_SOURCES and board >= 200:
        f[f"large company ({board} openings)"] = w["large_board"]
    elif job["source"] == "workday":
        f["big-company Workday portal"] = w["workday_portal"]

    # 5. A person I can actually write to
    job["contact"] = " | ".join(contacts_for(job, text))
    if "email in posting" in job["contact"]:
        f["contact email in posting"] = w["email_in_posting"]
    if "you know" in job["contact"]:
        f["you know someone there"] = w["connection_at_company"]
    if job["source"] not in FETCHERS:
        f["on public new-grad list"] = w["public_list"]
    if re.search(r"remote", loc, re.I):
        f["remote, national pool"] = w["remote"]
    elif HOME_RE.search(loc):
        f["Boston area"] = w["home"]
    elif WEST_ONLY_RE.search(loc) or WEST_ONLY_RE.search(text[:3000] if text else ""):
        f["West Coast hours/location"] = w["west_coast_only"]
    elif HUB_RE.search(loc):
        f["tech hub"] = w["hub"]
    else:
        f["smaller market"] = w["non_hub"]
    age = (datetime.now() - datetime.strptime(job["posted"], "%Y-%m-%d")).days if job.get("posted") else 7
    if age <= 2:
        f["just posted"] = w["fresh"]
    elif age > 14:
        f["posted 2+ weeks ago"] = w["stale"]
    elif age > 7:
        f["posted 1+ week ago"] = w["aging"]
    for pattern, label, mult in EDGE_RES:
        if pattern.search(text) or pattern.search(title):
            f[label] = mult

    score = 100.0
    for mult in f.values():
        score *= mult
    ups = [k for k, v in sorted(f.items(), key=lambda kv: -kv[1]) if v > 1.02][:3]
    downs = [k for k, v in sorted(f.items(), key=lambda kv: kv[1]) if v < 0.9][:3]
    why = "; ".join((["+ " + ", ".join(ups)] if ups else []) + (["- " + ", ".join(downs)] if downs else []))
    return max(1, min(99, round(score))), why


# --------------------------------------------------------------------------- output
RANK = {"Strong fit": 0, "Okay fit": 1, "Open the posting": 2, "Stretch": 3, "Unknown (no recognized skills)": 4}


def read_my_marks() -> dict[str, tuple]:
    """Keep whatever I typed in the Applied and Notes columns across runs (oldest file first, newest wins)."""
    from openpyxl import load_workbook
    marks = {}
    for path in sorted(HERE.glob("Jobs_*.xlsx"), key=lambda f: f.stat().st_mtime):
        if path.name.startswith("~$"):
            continue
        ws = load_workbook(path)["Apply"]
        head = [c.value for c in ws[1]]
        i_url, i_app, i_note = head.index("Link"), head.index("Applied"), head.index("Notes")
        for row in ws.iter_rows(min_row=2):
            url = row[i_url].hyperlink.target if row[i_url].hyperlink else row[i_url].value
            if url and (row[i_app].value or row[i_note].value):
                marks[url] = (row[i_app].value, row[i_note].value)
    return marks


def output_path() -> Path:
    """Excel leaves a ~$ lock file while the sheet is open; never write underneath an open workbook."""
    if (HERE / f"~${OUT_PATH.name}").exists():
        alt = HERE / f"Jobs_{TODAY}_{datetime.now():%H%M}.xlsx"
        print(f"{OUT_PATH.name} is open in Excel, so this run is saved as {alt.name}.")
        return alt
    return OUT_PATH


def write_xlsx(kept: list[dict], dropped: list[dict], marks: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    cols = ["Odds", "Why", "New", "Lane", "Company", "Title", "Location", "Posted", "Years asked", "Match",
            "Missing skills", "Visa wording in posting / pay", "Link", "Contact", "Find manager", "Applied", "Notes",
            "Source", "First seen"]

    def fill(ws, jobs):
        ws.append(cols)
        for j in jobs:
            applied, note = marks.get(j["url"], (None, None))
            lane_name, _, people = LANES[j["lane"]]
            ws.append([j["odds"], j["why"], "NEW" if j["first_seen"] == TODAY else "", lane_name,
                       j["company"], j["title"], j["location"][:80], j["posted"], j["years"],
                       j["score"] if j["score"] is not None else None, j["gaps"],
                       " || ".join(x for x in (j.get("visa_note"), j.get("sponsorship"), j.get("pay"), j.get("terms")) if x),
                       "open", j.get("contact", ""), "search", applied, note, j["source"], j["first_seen"]])
            row = ws.max_row
            cell = ws.cell(row=row, column=cols.index("Link") + 1)
            cell.hyperlink, cell.font = j["url"], Font(color="0563C1", underline="single")
            cell = ws.cell(row=row, column=cols.index("Find manager") + 1)
            query = urllib.parse.quote(f"{j['company']} {people}")
            cell.hyperlink = f"https://www.linkedin.com/search/results/people/?keywords={query}"
            cell.font = Font(color="0563C1", underline="single")
            if j["score"] is not None:
                ws.cell(row=row, column=cols.index("Match") + 1).number_format = "0%"
            if j["first_seen"] == TODAY:
                ws.cell(row=row, column=cols.index("New") + 1).fill = PatternFill("solid", fgColor="FFF2CC")

    ws = wb.active
    ws.title = "Apply"
    fill(ws, kept)

    ws2 = wb.create_sheet("Filtered out")
    ws2.append(["Reason", "Company", "Title", "Location", "Link", "Source"])
    for j in dropped:
        ws2.append([j["reason"], j["company"], j["title"], j["location"][:80], j["url"], j["source"]])

    apply_widths = [7, 48, 6, 30, 20, 46, 28, 11, 8, 8, 28, 60, 7, 40, 9, 9, 30, 20, 11]
    widths = {"Apply": apply_widths, "New today": apply_widths, "Filtered out": [34, 20, 50, 30, 60, 20]}
    for sheet in wb:
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for i, w in enumerate(widths[sheet.title], 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
    wb.save(out)
    return out


# --------------------------------------------------------------------------- commands
def run():
    profile_path = HERE.parent / "jdfit" / "profile.json"
    if not profile_path.exists():
        profile_path = profile_path.with_name("profile.example.json")
    profile = jdfit.load_json(profile_path)
    patterns = jdfit.compile_taxonomy(jdfit.load_json(HERE.parent / "jdfit" / "taxonomy.json"))
    companies = json.loads((HERE / "companies.json").read_text())
    tasks = [(FETCHERS[kind], item) for kind in FETCHERS for item in companies.get(kind, [])]
    tasks += [(github_list, entry) for entry in CONFIG["github_lists"]]

    failed = []

    def safe(task):
        fn, arg = task
        try:
            return fn(arg)
        except Exception as exc:  # one dead board should not stop the run
            failed.append(task)
            name = arg if isinstance(arg, str) else arg["name"]
            print(f"  ! {fn.__name__}:{name} failed: {exc}", file=sys.stderr)
            return []

    print(f"Fetching {len(tasks)} sources. This takes about 4 minutes; leave this window open.", flush=True)
    jobs, done = [], 0
    with ThreadPoolExecutor(max_workers=12) as pool:
        for batch in pool.map(safe, tasks):
            jobs.extend(batch)
            done += 1
            if done % 20 == 0:
                print(f"  {done} of {len(tasks)} sources done", flush=True)

    if len(failed) > len(tasks) * 0.15:
        sys.exit(f"{len(failed)} of {len(tasks)} sources failed, which looks like a network problem. "
                 "Nothing was written; the existing sheet is untouched. Run it again when the connection is back.")

    seen = json.loads(SEEN_PATH.read_text()) if SEEN_PATH.exists() else {}
    board_sizes = Counter((j["source"], j["company"]) for j in jobs)   # openings per company: a rough size signal
    kept, dropped, urls, already = [], [], set(), 0
    for job in jobs:
        job["location"] = job.get("location") or ""
        job["board_size"] = board_sizes[(job["source"], job["company"])]
        if job["url"] in urls:
            continue
        urls.add(job["url"])
        keep, reason = judge(job, profile, patterns)
        if keep:
            job["first_seen"] = seen.setdefault(job["url"], TODAY)
            if job["first_seen"] != TODAY:                  # already shown on an earlier day
                already += 1
                continue
            job["lane"] = lane_of(job["title"])
            job["odds"], job["why"] = odds(job, job.pop("_report"), job.pop("_text"))
            kept.append(job)
        else:
            job["reason"] = reason
            dropped.append(job)

    kept.sort(key=lambda j: j["posted"], reverse=True)           # newest first among equal scores
    kept.sort(key=lambda j: -j["odds"])
    counts = Counter(j["reason"].split(":")[0] for j in dropped)
    # Title rejections are the bulk and not worth auditing row by row; keep the substantive ones.
    dropped = sorted((j for j in dropped if not j["reason"].startswith(("title", "old"))),
                     key=lambda j: (j["reason"], j["company"]))
    out = write_xlsx(kept, dropped, read_my_marks(), output_path())
    SEEN_PATH.write_text(json.dumps(seen, indent=0))

    new = [j for j in kept if j["first_seen"] == TODAY]
    print(f"{len(jobs)} postings fetched from {len(tasks)} sources")
    print(f"{len(kept)} new jobs posted since {OLDEST} pass the filters "
          f"({already} others were already on an earlier sheet), {len(dropped)} filtered out for cause")
    for reason, n in counts.most_common():
        print(f"  filtered, {reason}: {n}")
    print(f"wrote {out.name}")
    (HERE / "last_output.txt").write_text(str(out))


def discover(path: str):
    """Try each name in a text file against all three job boards and print the ones that exist."""
    names = [n.strip().lower() for n in Path(path).read_text().split() if n.strip()]
    probes = {"greenhouse": "https://boards-api.greenhouse.io/v1/boards/{}/jobs",
              "lever": "https://api.lever.co/v0/postings/{}?mode=json&limit=1",
              "ashby": "https://api.ashbyhq.com/posting-api/job-board/{}"}

    def probe(pair):
        kind, name = pair
        try:
            data = get(probes[kind].format(name), timeout=20)
            count = len(data) if isinstance(data, list) else len(data.get("jobs", []))
            return (kind, name) if count else None
        except Exception:
            return None

    with ThreadPoolExecutor(max_workers=16) as pool:
        found = [r for r in pool.map(probe, [(k, n) for n in names for k in probes]) if r]
    result = {k: sorted(n for kk, n in found if kk == k) for k in probes}
    print(json.dumps(result, indent=1))


if __name__ == "__main__":
    if len(sys.argv) > 2 and sys.argv[1] == "discover":
        discover(sys.argv[2])
    elif len(sys.argv) > 1 and sys.argv[1] == "rescore":
        # Re-rank everything posted in the last N days under the current rules, without touching the daily
        # sweep's memory of what I've already seen:  jobfeed.py rescore 14
        days = int(sys.argv[2]) if len(sys.argv) > 2 else 14
        OLDEST = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d")
        SEEN_PATH = HERE / "rescore_seen.tmp.json"
        SEEN_PATH.unlink(missing_ok=True)
        OUT_PATH = HERE / f"Jobs_rescored_{days}days_{TODAY}.xlsx"
        last = (HERE / "last_output.txt").read_text() if (HERE / "last_output.txt").exists() else None
        run()
        SEEN_PATH.unlink(missing_ok=True)
        if last:
            (HERE / "last_output.txt").write_text(last)
    else:
        run()
