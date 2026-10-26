"""Turn job-alert emails (saved as text files in alerts_inbox/) into an "Alerts" sheet in today's job workbook.

The daily routine saves each new LinkedIn alert email's plain-text body as alerts_inbox/<message id>.txt,
then runs:  python alerts.py
Jobs already shown on an earlier day are skipped (alerts_seen.json), and processed files move to alerts_inbox/done/.
"""
import json
import re
from datetime import date
from pathlib import Path

HERE = Path(__file__).parent
INBOX = HERE / "alerts_inbox"
SEEN_PATH = HERE / "alerts_seen.json"
TODAY = date.today().isoformat()

# Titles that are clearly outside Jerry's lanes (retail, food, clinical, trades, driving, childcare, security guard work).
SKIP_TITLE = re.compile(
    r"nurse|\bRN\b|\bLPN\b|physician|pharmac|therap|dental|medical assistant|patient care|caregiver|clinical|"
    r"phlebotom|radiolog|sonograph|surgical|veterinar|behavior technician|\bRBT\b|BCBA|counselor|social worker|"
    r"cashier|barista|server|bartender|cook|chef|dishwasher|host\b|retail|store associate|sales associate|"
    r"merchandis|stocker|team member|crew member|key holder|beauty advisor|brand ambassador|"
    r"driver|delivery|warehouse|forklift|janitor|custodian|housekeep|laundry|environmental services|porter|"
    r"security officer|guard|soc operator|asset protection|loss prevention|customer service associate|lifeguard|swim|fitness|coach|teacher|tutor|nanny|babysit|childcare|"
    r"electrician|plumber|hvac|mechanic|technician - automotive|detailer|carpenter|painter|"
    r"sales representative|account executive|insurance agent|real estate agent", re.I)
SKIP_COMPANY = re.compile(r"walmart|target\b|macy|tjx|marshalls|home depot|lowe'?s|cvs|walgreens|whole foods|"
                          r"stop & shop|shaw'?s|trader joe|starbucks|securitas|allied universal|gardaworld|"
                          r"ymca|enterprise mobility|enterprise rent|ups\b|fedex|amazon (fresh|flex)", re.I)
WANT_LOCATION = re.compile(r"boston|brookline|cambridge|somerville|newton|watertown|waltham|quincy|medford|"
                           r"brighton|allston|malden|everett|chelsea|needham|dedham|massachusetts|\bMA\b|remote|united states", re.I)
TYPE_RE = re.compile(r"part[- ]?time|per diem|contract|temporary|temp\b|seasonal|freelance|hourly|\$\d+/hr", re.I)


def parse(body: str, source_id: str) -> list[dict]:
    """LinkedIn alert digests list each job as: title / company / location / ... / View job: <url>."""
    jobs = []
    search = re.search(r"job alert for (.+?) in ", body)
    alert_name = search.group(1) if search else ""
    for block in re.split(r"\n-{10,}\n", body):
        m = re.search(r"View job: (https://www\.linkedin\.com/comm/jobs/view/(\d+)/)", block)
        if not m:
            continue
        lines = [l.strip() for l in block[:m.start()].splitlines() if l.strip()]
        lines = [l for l in lines if not re.match(r"(Your job alert|\d+\+? new jobs|Apply with|\d+ (school alum|connection))", l)]
        if len(lines) < 3:
            continue
        title, company, location = lines[-3], lines[-2], lines[-1]
        jobs.append(dict(id=m.group(2), title=title, company=company, location=location,
                         url=f"https://www.linkedin.com/jobs/view/{m.group(2)}/", alert=alert_name, email=source_id))
    return jobs


def main():
    seen = json.loads(SEEN_PATH.read_text()) if SEEN_PATH.exists() else {}
    files = sorted(INBOX.glob("*.txt")) if INBOX.exists() else []
    found, skipped = [], 0
    for f in files:
        for j in parse(f.read_text(), f.stem):
            if j["id"] in seen:
                continue
            seen[j["id"]] = TODAY
            if SKIP_TITLE.search(j["title"]) or SKIP_COMPANY.search(j["company"]) or not WANT_LOCATION.search(j["location"]):
                skipped += 1
                continue
            j["type"] = (TYPE_RE.search(j["title"]).group(0).title() if TYPE_RE.search(j["title"]) else "")
            found.append(j)
    if not files:
        print("No alert emails in alerts_inbox/.")
        return

    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import Font, PatternFill
    target = Path((HERE / "last_output.txt").read_text().strip()) if (HERE / "last_output.txt").exists() else None
    if not target or not target.exists() or (target.parent / f"~${target.name}").exists():
        target = HERE / f"Alerts_{TODAY}.xlsx"
        wb = Workbook(); wb.remove(wb.active)
    else:
        wb = load_workbook(target)
    if "Alerts (LinkedIn)" in wb.sheetnames:
        ws = wb["Alerts (LinkedIn)"]
    else:
        ws = wb.create_sheet("Alerts (LinkedIn)", 1 if len(wb.sheetnames) > 0 else 0)
        ws.append(["Type", "Company", "Title", "Location", "Link", "From alert", "Applied", "Notes", "First seen"])
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
        for col, w in zip("ABCDEFGHI", [12, 28, 52, 26, 8, 34, 9, 30, 11]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    for j in sorted(found, key=lambda j: (j["type"] == "", j["company"])):
        ws.append([j["type"], j["company"], j["title"], j["location"], "open", j["alert"], None, None, TODAY])
        c = ws.cell(row=ws.max_row, column=5)
        c.hyperlink, c.font = j["url"], Font(color="0563C1", underline="single")
    wb.save(target)
    SEEN_PATH.write_text(json.dumps(seen, indent=0))
    done = INBOX / "done"
    done.mkdir(exist_ok=True)
    for f in files:
        f.rename(done / f.name)
    print(f"{len(found)} new alert jobs added to {target.name} ({skipped} skipped as retail/clinical/other areas).")


if __name__ == "__main__":
    main()
